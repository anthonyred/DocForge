"""Inspect paragraph and table placement in generated DOCX containers."""

from pathlib import Path
from zipfile import ZipFile
import json

from docx import Document
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

for path in (
    Path(r"D:\DocForge\.visual_validation\generated2.docx"),
    Path(r"D:\DocForge\.visual_validation\table_probe3.docx"),
):
    document = Document(path)
    print(path.name, "paragraphs", json.dumps([p.text for p in document.paragraphs], ensure_ascii=True))
    print("cells", json.dumps([[[c.text for c in row.cells] for row in table.rows] for table in document.tables], ensure_ascii=True))
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find(W + "body")
    print("body children", [item.tag.rsplit("}", 1)[-1] for item in body])
    print("table rows", [["".join(cell.itertext()) for cell in row.findall(W + "tc")] for table in body.findall(W + "tbl") for row in table.findall(W + "tr")])

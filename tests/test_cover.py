"""Unit tests for CoverPageBuilder."""

from pathlib import Path
import docx
from docx import Document
import pytest

from docforge.app.config import AppSettings
from docforge.engine.cover import CoverPageBuilder
from docforge.engine.profile import FormatProfile


def test_cover_page_executive_layout():
    """Verify cover page generates for executive_modern layout."""
    doc = Document()
    builder = CoverPageBuilder()
    profile = FormatProfile()
    profile.cover.layout = "executive_modern"
    settings = AppSettings(author_name="Anthony Reddy")

    builder.build_cover_page(doc, profile, settings, markdown_title="System Architecture Guide")

    # Document should have paragraphs and page break
    assert len(doc.paragraphs) > 0
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "System Architecture Guide" in full_text
    assert "Anthony Reddy" in full_text


def test_cover_page_disabled():
    """Verify cover page is skipped when disabled in profile."""
    doc = Document()
    builder = CoverPageBuilder()
    profile = FormatProfile()
    profile.cover.enabled = False
    settings = AppSettings()

    builder.build_cover_page(doc, profile, settings, markdown_title="Skipped Cover")
    assert len(doc.paragraphs) == 0

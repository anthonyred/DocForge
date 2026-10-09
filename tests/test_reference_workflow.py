"""Tests for arbitrary reference analysis and Pandoc-backed conversion."""

from pathlib import Path
import base64
import shutil
import subprocess
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import pytest

from docforge.engine.pandoc_renderer import PandocRenderer
from docforge.engine.reference_analyzer import ReferenceDocxAnalyzer
from docforge.engine.reference_mapper import prepare_reference_docx
from docforge.engine.reference_profile_store import ReferenceProfileStore


def _make_reference(path: Path, font: str, margin: float) -> None:
    """Create a small, distinct reference DOCX fixture."""
    doc = Document()
    doc.sections[0].left_margin = Inches(margin)
    doc.styles["Normal"].font.name = font
    doc.styles["Heading 1"].font.name = font
    doc.add_heading("Reference", level=1)
    doc.add_paragraph("Sample body")
    doc.save(path)


def test_analyzer_reports_distinct_reference_properties(tmp_path: Path):
    """Different references produce distinct analysis rather than hardcoded styles."""
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    _make_reference(first, "Arial", 1.0)
    _make_reference(second, "Times New Roman", 2.0)

    analyzer = ReferenceDocxAnalyzer()
    first_report, second_report = analyzer.analyze(first), analyzer.analyze(second)

    assert first_report.properties["Normal"].startswith("Arial")
    assert second_report.properties["Normal"].startswith("Times New Roman")
    assert first_report.properties["Margins"] != second_report.properties["Margins"]
    assert "--reference-doc" in first_report.strategy


def test_analyzer_rejects_corrupt_reference(tmp_path: Path):
    """Report a useful validation error for corrupt reference files."""
    path = tmp_path / "bad.docx"
    path.write_text("not a Word document", encoding="utf-8")
    with pytest.raises(ValueError, match="corrupt"):
        ReferenceDocxAnalyzer().analyze(path)


def test_analyzer_accepts_modern_ooxml_alignment_values(tmp_path: Path):
    """Describe alignment values that python-docx does not map to its enum."""
    path = tmp_path / "modern-alignment.docx"
    doc = Document()
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), "start")
    doc.styles["Normal"]._element.get_or_add_pPr().append(justification)
    doc.save(path)

    report = ReferenceDocxAnalyzer().analyze(path)

    assert "alignment=start" in report.properties["Normal"]


def test_analyzer_maps_custom_outline_style(tmp_path: Path):
    """Detect custom-named paragraph styles that declare a heading outline level."""
    path = tmp_path / "custom.docx"
    doc = Document()
    style = doc.styles.add_style("Executive Heading", WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    style._element.get_or_add_pPr().append(outline)
    doc.save(path)

    report = ReferenceDocxAnalyzer().analyze(path)

    assert report.mapped_styles["Markdown H1"] == "Executive Heading"


def test_reference_preparation_preserves_inherited_body_font(tmp_path: Path):
    """Copying Body Text must not clear font properties inherited from Normal."""
    source, prepared = tmp_path / "source.docx", tmp_path / "prepared.docx"
    _make_reference(source, "Arial", 1.0)
    source_document = Document(source)
    source_document.styles["Normal"].font.size = Pt(11)
    source_document.save(source)

    prepare_reference_docx(source, prepared)

    normal = Document(prepared).styles["Normal"].font
    assert normal.name == "Arial"
    assert normal.size.pt == 11


def test_reference_preparation_supplies_pandoc_compact_and_table_styles(tmp_path: Path):
    """References without Pandoc's built-in styles retain list and table layout."""
    source, prepared = tmp_path / "source.docx", tmp_path / "prepared.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).style = "Table Grid"
    doc.save(source)

    prepare_reference_docx(source, prepared)

    styles = Document(prepared).styles
    assert styles["Compact"].base_style.name == "Normal"
    assert styles["Compact"].paragraph_format.space_after.pt == 0
    assert styles["Table"].base_style.name in {"Light Shading", "Table Grid", "Normal Table"}


def test_reference_profile_store_saves_private_copy_and_versioned_report(tmp_path: Path):
    """Save reusable analysis separately while leaving its source unchanged."""
    source = tmp_path / "customer-template.docx"
    _make_reference(source, "Arial", 1.0)
    original = source.read_bytes()
    analysis = ReferenceDocxAnalyzer().analyze(source)

    result = ReferenceProfileStore(tmp_path / "profiles").save("Support Guide", analysis)

    assert (result / "reference.docx").read_bytes() == original
    assert json.loads((result / "profile.json").read_text(encoding="utf-8"))["profile_version"] == 1
    assert ReferenceProfileStore(tmp_path / "profiles").list_profiles() == [
        ("Support Guide", result / "reference.docx")
    ]
    assert source.read_bytes() == original


def test_pandoc_renderer_uses_reference_and_preserves_source(tmp_path: Path, monkeypatch):
    """Invoke Pandoc with safe arguments and keep the reference untouched."""
    reference = tmp_path / "reference.docx"
    _make_reference(reference, "Arial", 1.0)
    original_bytes = reference.read_bytes()
    output = tmp_path / "result.docx"
    captured = {}

    def fake_run(command, **kwargs):
        captured["command"] = command
        captured["markdown"] = Path(command[4]).read_text(encoding="utf-8")
        captured["reference"] = Path(command[3].split("=", 1)[1])
        captured["prepared_h1_font"] = Document(captured["reference"]).styles["Heading 1"].font.name
        fake_output = Document()
        fake_output.add_paragraph("Pandoc test output")
        fake_output.save(command[-1])
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("docforge.engine.pandoc_renderer.subprocess.run", fake_run)
    result = PandocRenderer(executable="pandoc-test").render("# Hello", reference, output)

    assert result == output.resolve()
    assert captured["markdown"] == "# Hello"
    assert captured["reference"] != reference.resolve()
    assert captured["prepared_h1_font"] == "Arial"
    assert reference.read_bytes() == original_bytes


def test_pandoc_renderer_rejects_reference_overwrite(tmp_path: Path):
    """Never allow conversion to overwrite the selected reference file."""
    reference = tmp_path / "reference.docx"
    _make_reference(reference, "Arial", 1.0)
    with pytest.raises(ValueError, match="must not overwrite"):
        PandocRenderer(executable="pandoc-test").render("# Hello", reference, reference)


def test_pandoc_renderer_rejects_existing_output_without_confirmation(tmp_path: Path):
    """Avoid silently replacing an existing output from headless callers."""
    reference = tmp_path / "reference.docx"
    output = tmp_path / "existing.docx"
    _make_reference(reference, "Arial", 1.0)
    output.write_bytes(b"keep me")

    with pytest.raises(FileExistsError, match="already exists"):
        PandocRenderer(executable="pandoc-test").render("# Hello", reference, output)

    assert output.read_bytes() == b"keep me"


def test_pandoc_renderer_reports_missing_binary(tmp_path: Path, monkeypatch):
    """Give a direct install or PATH error when Pandoc is unavailable."""
    reference = tmp_path / "reference.docx"
    _make_reference(reference, "Arial", 1.0)
    monkeypatch.setattr("docforge.engine.pandoc_renderer.shutil.which", lambda _: None)
    with pytest.raises(FileNotFoundError, match="Install Pandoc"):
        PandocRenderer().render("# Hello", reference, tmp_path / "output.docx")


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_real_pandoc_generates_editable_docx(tmp_path: Path):
    """Smoke-test the installed Pandoc binary with a generated reference."""
    reference = tmp_path / "reference.docx"
    output = tmp_path / "output.docx"
    _make_reference(reference, "Arial", 1.0)
    PandocRenderer().render("# Report\n\nGenerated with Pandoc.", reference, output)

    assert "Generated with Pandoc." in "\n".join(p.text for p in Document(output).paragraphs)


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_real_pandoc_retains_distinct_reference_styles(tmp_path: Path):
    """Different references yield different heading styles in generated DOCX files."""
    first_ref, second_ref = tmp_path / "first.docx", tmp_path / "second.docx"
    first_out, second_out = tmp_path / "first-out.docx", tmp_path / "second-out.docx"
    _make_reference(first_ref, "Arial", 1.0)
    _make_reference(second_ref, "Times New Roman", 2.0)

    renderer = PandocRenderer()
    renderer.render("# Report", first_ref, first_out)
    renderer.render("# Report", second_ref, second_out)

    first_font = Document(first_out).styles["Heading 1"].font.name
    second_font = Document(second_out).styles["Heading 1"].font.name
    assert first_font == "Arial"
    assert second_font == "Times New Roman"


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_pandoc_preserves_markdown_structures_unicode_and_citation(tmp_path: Path):
    """Verify Pandoc keeps representative content in editable DOCX structure."""
    reference = tmp_path / "reference.docx"
    output = tmp_path / "structures.docx"
    _make_reference(reference, "Arial", 1.0)
    markdown = """# Café report

Unicode — 日本語, **bold**, *italic*, and [a link](https://example.com).

1. Parent
   - Nested item

| Name | State |
| --- | --- |
| Worker | Ready |

```python
print('hello')
```

Citation [@doe2020].
"""

    PandocRenderer().render(markdown, reference, output)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Café report" in text
    assert "日本語" in text
    assert "Nested item" in text
    assert "print('hello')" in text
    assert "doe2020" in text
    assert len(document.tables) == 1


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_pandoc_maps_custom_heading_style_to_markdown_heading(tmp_path: Path):
    """Apply an observed custom outline style to Pandoc's heading mapping."""
    reference, output = tmp_path / "custom.docx", tmp_path / "custom-output.docx"
    doc = Document()
    style = doc.styles.add_style("Executive Heading", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Trebuchet MS"
    style.font.size = Pt(31)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    style._element.get_or_add_pPr().append(outline)
    doc.save(reference)

    PandocRenderer().render("# Executive summary", reference, output)

    heading_style = Document(output).styles["Heading 1"]
    assert heading_style.font.name == "Trebuchet MS"
    assert heading_style.font.size.pt == 31


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_pandoc_maps_custom_code_style(tmp_path: Path):
    """Apply a clearly named custom code style to Pandoc code paragraphs."""
    reference, output = tmp_path / "custom-code.docx", tmp_path / "custom-code-output.docx"
    doc = Document()
    style = doc.styles.add_style("Terminal Code", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Consolas"
    style.font.size = Pt(8.5)
    doc.save(reference)

    PandocRenderer().render("```text\ncommand\n```", reference, output)

    code_style = Document(output).styles["Source Code"]
    assert code_style.font.name == "Consolas"
    assert code_style.font.size.pt == 8.5


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_pandoc_resolves_images_relative_to_markdown_source(tmp_path: Path):
    """Resolve local image links from the selected Markdown file directory."""
    reference, output = tmp_path / "reference.docx", tmp_path / "image-output.docx"
    _make_reference(reference, "Arial", 1.0)
    image_path = tmp_path / "pixel.png"
    image_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p8sAAAAASUVORK5CYII="
    ))

    PandocRenderer().render("![pixel](pixel.png)", reference, output, resource_paths=[tmp_path])

    assert len(Document(output).inline_shapes) == 1


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_pandoc_retains_reference_page_geometry_and_header_footer(tmp_path: Path):
    """Retain document properties and header/footer from the selected reference."""
    reference, output = tmp_path / "layout.docx", tmp_path / "layout-output.docx"
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1.75)
    section.header.paragraphs[0].text = "Reference Header"
    section.footer.paragraphs[0].text = "Reference Footer"
    doc.save(reference)

    PandocRenderer().render("# Layout check", reference, output)

    generated = Document(output)
    assert abs(generated.sections[0].left_margin / Inches(1) - 1.75) < 0.02
    assert generated.sections[0].header.paragraphs[0].text == "Reference Header"
    assert generated.sections[0].footer.paragraphs[0].text == "Reference Footer"
